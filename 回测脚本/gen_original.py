"""Original strategy: Volty EMA50 + TP1.5%, no RSI, no dynamic sizing"""
import pandas as pd, numpy as np, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INIT=700; PCT=60; LEV=3; SL=5; TP=1.5

df=pd.read_csv("/opt/trading/ETHUSDT_15m_full.csv",parse_dates=["ts"],index_col="ts")
n=len(df)
H=df["high"]; L=df["low"]; C=df["close"]; O=df["open"]
tr=pd.concat([H-L,(H-C.shift(1)).abs(),(L-C.shift(1)).abs()],axis=1).max(axis=1)
atr5=tr.rolling(5).mean()*0.75; ema50=C.ewm(span=50,adjust=False).mean()

bal=INIT; pos=None; ep=uv=0.0; trades=[]; tp_c=0; sl_c=0
for i in range(200,n):
    hi,lo,op,cl=H.iloc[i],L.iloc[i],O.iloc[i],C.iloc[i]; tm=df.index[i]
    if pos=="LONG":
        sl_p=ep*(1-SL/100)
        if lo<=sl_p:
            pnl=(sl_p-ep)/ep*uv; bal+=pnl; sl_c+=1
            trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"LONG","止损",round(ep,2),round(sl_p,2),round(uv,2),round(pnl,2),round(bal,2),tm.year]); pos=None; continue
    if pos=="SHORT":
        sl_p=ep*(1+SL/100)
        if hi>=sl_p:
            pnl=(ep-sl_p)/ep*uv; bal+=pnl; sl_c+=1
            trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"SHORT","止损",round(ep,2),round(sl_p,2),round(uv,2),round(pnl,2),round(bal,2),tm.year]); pos=None; continue
    if pos=="LONG" and hi>=ep*(1+TP/100):
        tp_p=ep*(1+TP/100); pnl=(tp_p-ep)/ep*uv; bal+=pnl; tp_c+=1
        trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"LONG","止盈",round(ep,2),round(tp_p,2),round(uv,2),round(pnl,2),round(bal,2),tm.year]); pos=None; continue
    if pos=="SHORT" and lo<=ep*(1-TP/100):
        tp_p=ep*(1-TP/100); pnl=(ep-tp_p)/ep*uv; bal+=pnl; tp_c+=1
        trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"SHORT","止盈",round(ep,2),round(tp_p,2),round(uv,2),round(pnl,2),round(bal,2),tm.year]); pos=None; continue
    pc=C.iloc[i-1]; pa=atr5.iloc[i-1]; pe=ema50.iloc[i-1]
    if pd.isna(pa) or pa<=0: continue
    ll_v=pc+pa; sl_v=pc-pa
    lt=hi>=ll_v and pc>pe and pos!="LONG"
    st=lo<=sl_v and pc<pe and pos!="SHORT"
    if lt and st:
        if abs(op-ll_v)>abs(op-sl_v): lt=False
        else: st=False
    if lt:
        if pos=="SHORT": xp=ll_v; pl=(ep-xp)/ep*uv; bal+=pl; trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"SHORT->LONG","翻转",round(ep,2),round(xp,2),round(uv,2),round(pl,2),round(bal,2),tm.year]); pos=None
        if bal>10: cap=bal*PCT/100*0.98; uv=cap*LEV; ep=ll_v; pos="LONG"
    elif st:
        if pos=="LONG": xp=sl_v; pl=(xp-ep)/ep*uv; bal+=pl; trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"LONG->SHORT","翻转",round(ep,2),round(xp,2),round(uv,2),round(pl,2),round(bal,2),tm.year]); pos=None
        if bal>10: cap=bal*PCT/100*0.98; uv=cap*LEV; ep=sl_v; pos="SHORT"
if pos:
    lp=C.iloc[-1]; pl=(lp-ep)/ep*uv if pos=="LONG" else (ep-lp)/ep*uv; bal+=pl
    trades.append([len(trades)+1,str(df.index[-1].date()),str(df.index[-1].time())[:5],pos,"最终",round(ep,2),round(lp,2),round(uv,2),round(pl,2),round(bal,2),df.index[-1].year])

final_bal=round(bal,2); total_ret=round((final_bal-INIT)/INIT*100,2)
closed=[t for t in trades if t[8]!=0]
wins=[t for t in closed if t[8]>0]; losses=[t for t in closed if t[8]<0]
wr=len(wins)/len(closed)*100 if closed else 0

yearly=[]
for y in range(2019,2027):
    yt=[t for t in closed if t[10]==y]
    if not yt: continue
    ypnl=sum(t[8] for t in yt)
    prev=[t for t in closed if t[10]<y]; ys=prev[-1][9] if prev else INIT; ye=yt[-1][9]
    yearly.append([y,len(yt),len([t for t in yt if t[8]>0]),len([t for t in yt if t[8]<0]),round(ypnl,2),round((ye-ys)/ys*100,2) if ys>0 else 0,round(ys,2),round(ye,2)])

# Excel
wb=Workbook()
hfont=Font(name="Microsoft YaHei",size=11,bold=True,color="FFFFFF")
hfill=PatternFill(start_color="1a1a2e",end_color="1a1a2e",fill_type="solid")
tfont=Font(name="Microsoft YaHei",size=16,bold=True,color="1a1a2e")
sf=Font(name="Microsoft YaHei",size=10,color="666666")
dfont=Font(name="Microsoft YaHei",size=10)
gf=Font(name="Microsoft YaHei",size=10,color="006100"); rf_=Font(name="Microsoft YaHei",size=10,color="9C0006")
gfill=PatternFill(start_color="C6EFCE",end_color="C6EFCE",fill_type="solid")
rfill=PatternFill(start_color="FFC7CE",end_color="FFC7CE",fill_type="solid")
bdr=Border(left=Side(style="thin",color="D0D0D0"),right=Side(style="thin",color="D0D0D0"),top=Side(style="thin",color="D0D0D0"),bottom=Side(style="thin",color="D0D0D0"))
lfill=PatternFill(start_color="F5F5F5",end_color="F5F5F5",fill_type="solid")

def hdr(ws,row,cols):
    for c in range(1,cols+1): cl=ws.cell(row=row,column=c); cl.font=hfont; cl.fill=hfill; cl.alignment=Alignment(horizontal="center",vertical="center"); cl.border=bdr
def sdr(ws,row,cols,even=False):
    for c in range(1,cols+1): cl=ws.cell(row=row,column=c); cl.font=dfont; cl.border=bdr; cl.alignment=Alignment(horizontal="center")
    if even:
        for c in range(1,cols+1): ws.cell(row=row,column=c).fill=lfill

ws1=wb.active; ws1.title="每笔交易"
ws1.merge_cells("A1:K1"); ws1.cell(row=1,column=1,value="原始策略回测 Volty EMA50 TP1.5%").font=tfont
ws1.merge_cells("A2:K2"); ws1.cell(row=2,column=1,value="币安ETHUSDT 15m {} {} {}%x{}x SL{}% TP{}%".format(str(df.index[0])[:10],str(df.index[-1])[:10],PCT,LEV,SL,TP)).font=sf
h1=["序号","日期","时间","方向","动作","入场价","出场价","持仓价值","盈亏","余额","年份"]
for c,h in enumerate(h1,1): ws1.cell(row=4,column=c,value=h); hdr(ws1,4,len(h1))
for i in range(len(trades)):
    r=i+5; t=trades[i]
    for c in range(len(h1)): ws1.cell(row=r,column=c+1,value=t[c] if c<10 else t[c])
    sdr(ws1,r,len(h1),i%2==0)
    pc_=ws1.cell(row=r,column=9)
    if t[8]>0: pc_.font=gf; pc_.fill=gfill
    elif t[8]<0: pc_.font=rf_; pc_.fill=rfill
ws1.freeze_panes="A5"
for c in range(1,12): ws1.column_dimensions[get_column_letter(c)].width=14

ws2=wb.create_sheet("年度汇总")
ws2.merge_cells("A1:H1"); ws2.cell(row=1,column=1,value="年度收益汇总").font=tfont
h2=["年份","交易次数","盈利次数","亏损次数","盈亏","收益率%","年初余额","年末余额"]
for c,h in enumerate(h2,1): ws2.cell(row=3,column=c,value=h); hdr(ws2,3,len(h2))
for i,yr in enumerate(yearly):
    r=i+4
    for c in range(len(h2)): ws2.cell(row=r,column=c+1,value=yr[c]); sdr(ws2,r,len(h2),i%2==0)
    rc_=ws2.cell(row=r,column=6)
    if yr[5]>0: rc_.font=gf; rc_.fill=gfill
    elif yr[5]<0: rc_.font=rf_; rc_.fill=rfill
for c in range(1,9): ws2.column_dimensions[get_column_letter(c)].width=16; ws2.freeze_panes="A4"

out="/opt/trading/original_report.xlsx"; wb.save(out)
print("Saved:",out,"| Final: ${:,.0f} ({:+.1f}%)".format(final_bal,total_ret))
for y in yearly: print("  {}: ${:,.0f} -> ${:,.0f} ({:+.1f}%)".format(y[0],y[6],y[7],y[5]))
