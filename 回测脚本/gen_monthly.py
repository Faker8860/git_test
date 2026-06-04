"""Generate monthly PnL comparison Excel"""
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

INIT=700; PCT=60; LEV=3
df=pd.read_csv("/opt/trading/ETHUSDT_15m_full.csv",parse_dates=["ts"],index_col="ts")
n=len(df)
H=df["high"]; L=df["low"]; C=df["close"]; O=df["open"]
tr=pd.concat([H-L,(H-C.shift(1)).abs(),(L-C.shift(1)).abs()],axis=1).max(axis=1)
atr5=tr.rolling(5).mean()*0.75; ema50=C.ewm(span=50,adjust=False).mean()
atr14=tr.rolling(14).mean(); mid_atr=atr14.median()
d=C.diff(); g=d.clip(lower=0); l=(-d).clip(lower=0)
rsi=100-(100/(1+g.rolling(14).mean()/l.rolling(14).mean()))

def run(name,tp,use_rsi=False,use_dyn=False):
    bal=INIT; pos=None; ep=uv=0.0; eq={str(df.index[199])[:10]:INIT}
    for i in range(200,n):
        hi,lo,cl=H.iloc[i],L.iloc[i],C.iloc[i]
        if pos=="LONG" and lo<=ep*0.95: bal+=(ep*0.95-ep)/ep*uv; pos=None
        elif pos=="SHORT" and hi>=ep*1.05: bal+=(ep-ep*1.05)/ep*uv; pos=None
        elif pos=="LONG" and hi>=ep*(1+tp/100): bal+=(ep*(1+tp/100)-ep)/ep*uv; pos=None
        elif pos=="SHORT" and lo<=ep*(1-tp/100): bal+=(ep-ep*(1-tp/100))/ep*uv; pos=None
        else:
            pc=C.iloc[i-1]; pa=atr5.iloc[i-1]; pe=ema50.iloc[i-1]
            if not pd.isna(pa) and pa>0:
                rb=rsi.iloc[i-1]>50 if (use_rsi and not pd.isna(rsi.iloc[i-1])) else True
                rbe=rsi.iloc[i-1]<50 if (use_rsi and not pd.isna(rsi.iloc[i-1])) else True
                lt=hi>=pc+pa and pc>pe and rb and pos!="LONG"
                st=lo<=pc-pa and pc<pe and rbe and pos!="SHORT"
                if lt and st:
                    if abs(O.iloc[i]-(pc+pa))>abs(O.iloc[i]-(pc-pa)): lt=False
                    else: st=False
                if lt:
                    if pos=="SHORT": bal+=(ep-(pc+pa))/ep*uv; pos=None
                    if bal>10:
                        cap=bal*PCT/100
                        if use_dyn and not pd.isna(atr14.iloc[i]): av=atr14.iloc[i]; cap*=min(1.5,max(0.5,mid_atr/av)) if av>0 else 1
                        uv=min(cap,bal*0.98)*LEV; ep=pc+pa; pos="LONG"
                elif st:
                    if pos=="LONG": bal+=((pc-pa)-ep)/ep*uv; pos=None
                    if bal>10:
                        cap=bal*PCT/100
                        if use_dyn and not pd.isna(atr14.iloc[i]): av=atr14.iloc[i]; cap*=min(1.5,max(0.5,mid_atr/av)) if av>0 else 1
                        uv=min(cap,bal*0.98)*LEV; ep=pc-pa; pos="SHORT"
        eq[str(df.index[i])[:10]]=bal
    if pos: bal+=((C.iloc[-1]-ep)/ep*uv if pos=="LONG" else (ep-C.iloc[-1])/ep*uv); eq[str(df.index[-1])[:10]]=bal
    # Monthly from daily last value
    daily={}
    for d,v in sorted(eq.items()):
        daily[d[:7]]=v  # keeps last value of each month
    return daily,bal

m1,b1=run("original",1.5); m2,b2=run("optimized",7,True,True)

wb=Workbook()
ws=wb.active; ws.title="Monthly PnL"
gfill=PatternFill(start_color="C6EFCE",end_color="C6EFCE",fill_type="solid")
rfill=PatternFill(start_color="FFC7CE",end_color="FFC7CE",fill_type="solid")
bdr=Border(left=Side(style="thin"),right=Side(style="thin"),top=Side(style="thin"),bottom=Side(style="thin"))
dfont=Font(name="Microsoft YaHei",size=10)
hfont=Font(name="Microsoft YaHei",size=11,bold=True)
gf=Font(name="Microsoft YaHei",size=10,color="006100")
rf_=Font(name="Microsoft YaHei",size=10,color="9C0006")

ws.merge_cells("A1:I1"); ws.cell(row=1,column=1,value="Monthly PnL: Original vs Optimized").font=Font(name="Microsoft YaHei",size=14,bold=True)
headers=["Month","Orig Start","Orig End","Orig PnL","Orig %","Opt Start","Opt End","Opt PnL","Opt %"]
for c,h in enumerate(headers,1): ws.cell(row=3,column=c,value=h).font=hfont
for c in range(1,10): ws.cell(row=3,column=c).border=bdr

row=4; last_o=INIT; last_p=INIT
for y in range(2019,2027):
    for m in range(1,13):
        label="{}-{:02d}".format(y,m)
        os_=last_o; oe_=last_o; opnl=0; oret=0
        if label in m1:
            oe_=m1[label]; opnl=round(oe_-os_,2); oret=round((oe_-os_)/os_*100,2) if os_>0 else 0; last_o=oe_
        ps_=last_p; pe_=last_p; ppnl=0; pret=0
        if label in m2:
            pe_=m2[label]; ppnl=round(pe_-ps_,2); pret=round((pe_-ps_)/ps_*100,2) if ps_>0 else 0; last_p=pe_

        ws.cell(row=row,column=1,value=label).font=dfont
        ws.cell(row=row,column=2,value=round(os_,0)).font=dfont
        ws.cell(row=row,column=3,value=round(oe_,0)).font=dfont
        c4=ws.cell(row=row,column=4,value=opnl); c4.font=gf if opnl>=0 else rf_
        c5=ws.cell(row=row,column=5,value=oret); c5.font=gf if oret>=0 else rf_
        ws.cell(row=row,column=6,value=round(ps_,0)).font=dfont
        ws.cell(row=row,column=7,value=round(pe_,0)).font=dfont
        c8=ws.cell(row=row,column=8,value=ppnl); c8.font=gf if ppnl>=0 else rf_
        c9=ws.cell(row=row,column=9,value=pret); c9.font=gf if pret>=0 else rf_

        if oret>0: c4.fill=gfill; c5.fill=gfill
        elif oret<0: c4.fill=rfill; c5.fill=rfill
        if pret>0: c8.fill=gfill; c9.fill=gfill
        elif pret<0: c8.fill=rfill; c9.fill=rfill

        for c in range(1,10): ws.cell(row=row,column=c).border=bdr
        row+=1

for c in range(1,10): ws.column_dimensions[chr(64+c)].width=14
ws.freeze_panes="A4"
wb.save("/opt/trading/monthly_compare.xlsx")
print("Done: {} months, Orig ${:,.0f}, Opt ${:,.0f}".format(row-4,last_o,last_p))
