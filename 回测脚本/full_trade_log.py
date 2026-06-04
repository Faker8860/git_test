"""Complete trade log: every entry & exit from $900"""
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INIT=900; PCT=60; LEV=3; SL=5; TP=7

df=pd.read_csv("/opt/trading/ETHUSDT_15m_full.csv",parse_dates=["ts"],index_col="ts")
n=len(df)
H=df["high"].values; L=df["low"].values; C=df["close"].values; O=df["open"].values
tr=pd.concat([df["high"]-df["low"],(df["high"]-df["close"].shift(1)).abs(),(df["low"]-df["close"].shift(1)).abs()],axis=1).max(axis=1)
atr5_v=tr.rolling(5).mean()*0.75; ema50_v=df["close"].ewm(span=50,adjust=False).mean()
atr14_v=tr.rolling(14).mean(); mid_atr=atr14_v.median()
d=df["close"].diff(); g=d.clip(lower=0); l=(-d).clip(lower=0)
rsi_v=(100-(100/(1+g.rolling(14).mean()/l.rolling(14).mean()))).values

bal=INIT; pos=None; ep=uv=0.0; trades=[]
for i in range(200,n):
    hi,lo,cl=H[i],L[i],C[i]; tm=df.index[i]
    # SL
    if pos==1 and lo<=ep*0.95:
        slp=ep*0.95; pnl=(slp-ep)/ep*uv; bal+=pnl
        trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"LONG","止损",round(ep,2),round(slp,2),round(uv,2),round(pnl,2),round(bal,2),tm.year,"SL"])
        pos=None; continue
    if pos==-1 and hi>=ep*1.05:
        slp=ep*1.05; pnl=(ep-slp)/ep*uv; bal+=pnl
        trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"SHORT","止损",round(ep,2),round(slp,2),round(uv,2),round(pnl,2),round(bal,2),tm.year,"SL"])
        pos=None; continue
    # TP
    if pos==1 and hi>=ep*1.07:
        tpp=ep*1.07; pnl=(tpp-ep)/ep*uv; bal+=pnl
        trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"LONG","止盈",round(ep,2),round(tpp,2),round(uv,2),round(pnl,2),round(bal,2),tm.year,"TP"])
        pos=None; continue
    if pos==-1 and lo<=ep*0.93:
        tpp=ep*0.93; pnl=(ep-tpp)/ep*uv; bal+=pnl
        trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"SHORT","止盈",round(ep,2),round(tpp,2),round(uv,2),round(pnl,2),round(bal,2),tm.year,"TP"])
        pos=None; continue
    # Signal
    pc=C[i-1]; pa=atr5_v.iloc[i-1]; pe=ema50_v.iloc[i-1]
    if pd.isna(pa) or pa<=0: continue
    rb=rsi_v[i-1]>50 if not pd.isna(rsi_v[i-1]) else True
    rbe=rsi_v[i-1]<50 if not pd.isna(rsi_v[i-1]) else True
    lt=hi>=pc+pa and pc>pe and rb and pos!=1
    st=lo<=pc-pa and pc<pe and rbe and pos!=-1
    if lt and st:
        if abs(O[i]-(pc+pa))>abs(O[i]-(pc-pa)): lt=False
        else: st=False
    if lt:
        if pos==-1:
            xp=pc+pa; pl=(ep-xp)/ep*uv; bal+=pl
            trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"SHORT->LONG","翻转",round(ep,2),round(xp,2),round(uv,2),round(pl,2),round(bal,2),tm.year,"FLIP"])
            pos=None
        if bal>10:
            cap=bal*PCT/100; av=atr14_v.iloc[i] if not pd.isna(atr14_v.iloc[i]) else mid_atr
            if av>0: cap*=min(1.2,max(0.5,mid_atr/av))
            uv=min(cap,bal*0.98)*LEV; ep=pc+pa; pos=1
            trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"LONG","开仓",round(ep,2),0,round(uv,2),0,round(bal,2),tm.year,"ENTRY"])
    elif st:
        if pos==1:
            xp=pc-pa; pl=(xp-ep)/ep*uv; bal+=pl
            trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"LONG->SHORT","翻转",round(ep,2),round(xp,2),round(uv,2),round(pl,2),round(bal,2),tm.year,"FLIP"])
            pos=None
        if bal>10:
            cap=bal*PCT/100; av=atr14_v.iloc[i] if not pd.isna(atr14_v.iloc[i]) else mid_atr
            if av>0: cap*=min(1.2,max(0.5,mid_atr/av))
            uv=min(cap,bal*0.98)*LEV; ep=pc-pa; pos=-1
            trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"SHORT","开仓",round(ep,2),0,round(uv,2),0,round(bal,2),tm.year,"ENTRY"])
if pos:
    lp=C[-1]; pl=(lp-ep)/ep*uv if pos==1 else (ep-lp)/ep*uv; bal+=pl
    trades.append([len(trades)+1,str(df.index[-1].date()),str(df.index[-1].time())[:5],pos,"最终平仓",round(ep,2),round(lp,2),round(uv,2),round(pl,2),round(bal,2),df.index[-1].year,"CLOSE"])

final_bal=round(bal,2); total_ret=round((final_bal-INIT)/INIT*100,2)
closed=[t for t in trades if t[8]!=0]
wins=[t for t in closed if t[8]>0]; losses=[t for t in closed if t[8]<0]
wr=len(wins)/len(closed)*100 if closed else 0

# Excel
wb=Workbook()
hfont=Font(name="Microsoft YaHei",size=11,bold=True,color="FFFFFF")
hfill=PatternFill(start_color="1a1a2e",end_color="1a1a2e",fill_type="solid")
tfont=Font(name="Microsoft YaHei",size=16,bold=True,color="1a1a2e")
sf=Font(name="Microsoft YaHei",size=10,color="666666")
dfont=Font(name="Microsoft YaHei",size=10)
gf=Font(name="Microsoft YaHei",size=10,color="006100"); rf_=Font(name="Microsoft YaHei",size=10,color="9C0006")
gfill=PatternFill(start_color="C6EFCE",end_color="C6EFCE",fill_type="solid")
rfill=PatternFill(start_color="FFC7CE",end_color="FFC7CE",fill_type="solid")
bdr=Border(left=Side(style="thin"),right=Side(style="thin"),top=Side(style="thin"),bottom=Side(style="thin"))
lfill=PatternFill(start_color="F5F5F5",end_color="F5F5F5",fill_type="solid")

def hdr(ws,row,cols):
    for c in range(1,cols+1): cl=ws.cell(row=row,column=c); cl.font=hfont; cl.fill=hfill; cl.alignment=Alignment(horizontal="center",vertical="center"); cl.border=bdr
def sdr(ws,row,cols,even):
    for c in range(1,cols+1): cl=ws.cell(row=row,column=c); cl.font=dfont; cl.border=bdr; cl.alignment=Alignment(horizontal="center")
    if even:
        for c in range(1,cols+1): ws.cell(row=row,column=c).fill=lfill

# Sheet 1: All trades
ws1=wb.active; ws1.title="Trades"
ws1.merge_cells("A1:M1"); ws1.cell(row=1,column=1,value="Current Strategy - Every Trade Log - Start ${}".format(INIT)).font=tfont
ws1.merge_cells("A2:M2"); ws1.cell(row=2,column=1,value="Volty+RSI+DynSize(1.2x) {}%x{}x SL{}% TP{}% | {} candles | {} -> {}".format(PCT,LEV,SL,TP,n,str(df.index[0])[:10],str(df.index[-1])[:10])).font=sf
h1=["#","Date","Time","Direction","Action","Entry","Exit","Value","PnL","Balance","Year","Type","Note"]
for c,h in enumerate(h1,1): ws1.cell(row=4,column=c,value=h)
hdr(ws1,4,len(h1))
for i,t in enumerate(trades):
    r=i+5
    # Pad trade to 13 columns
    padded = list(t) + [''] * (len(h1) - len(t))
    for c in range(len(h1)): ws1.cell(row=r,column=c+1,value=padded[c])
    sdr(ws1,r,len(h1),i%2==0)
    pnl_val = t[8] if len(t)>8 else 0
    if isinstance(pnl_val,(int,float)) and pnl_val>0: ws1.cell(row=r,column=9).font=gf; ws1.cell(row=r,column=9).fill=gfill
    elif isinstance(pnl_val,(int,float)) and pnl_val<0: ws1.cell(row=r,column=9).font=rf_; ws1.cell(row=r,column=9).fill=rfill
ws1.freeze_panes="A5"
for c in range(1,14): ws1.column_dimensions[get_column_letter(c)].width=13

# Sheet 2: Monthly
ws2=wb.create_sheet("Monthly")
ws2.merge_cells("A1:F1"); ws2.cell(row=1,column=1,value="Monthly Summary").font=tfont
h2=["Month","Trades","Wins","Losses","PnL","Return%"]
for c,h in enumerate(h2,1): ws2.cell(row=3,column=c,value=h); hdr(ws2,3,len(h2))
sr=4
for y in range(2019,2027):
    for m in range(1,13):
        mt=[t for t in closed if t[10]==y and t[1][5:7]=="{:02d}".format(m)]
        if not mt: continue
        label="{}/{:02d}".format(y,m); mpnl=sum(t[8] for t in mt)
        prev=[t for t in closed if (t[10]<y) or (t[10]==y and t[1][5:7]<"{:02d}".format(m))]
        ms=prev[-1][9] if prev else INIT; me=mt[-1][9]; mr=round((me-ms)/ms*100,2) if ms>0 else 0
        vals=[label,len(mt),len([t for t in mt if t[8]>0]),len([t for t in mt if t[8]<0]),round(mpnl,2),mr]
        for c,v in enumerate(vals,1): ws2.cell(row=sr,column=c,value=v); ws2.cell(row=sr,column=c).font=dfont; ws2.cell(row=sr,column=c).border=bdr; ws2.cell(row=sr,column=c).alignment=Alignment(horizontal="center")
        if mr>0: ws2.cell(row=sr,column=5).font=gf; ws2.cell(row=sr,column=5).fill=gfill; ws2.cell(row=sr,column=6).font=gf; ws2.cell(row=sr,column=6).fill=gfill
        elif mr<0: ws2.cell(row=sr,column=5).font=rf_; ws2.cell(row=sr,column=5).fill=rfill; ws2.cell(row=sr,column=6).font=rf_; ws2.cell(row=sr,column=6).fill=rfill
        sr+=1
for c in range(1,7): ws2.column_dimensions[get_column_letter(c)].width=14
ws2.freeze_panes="A4"

# Sheet 3: Yearly
ws3=wb.create_sheet("Yearly")
ws3.merge_cells("A1:F1"); ws3.cell(row=1,column=1,value="Yearly Summary").font=tfont
h3=["Year","Trades","Wins","Losses","PnL","Return%"]
for c,h in enumerate(h3,1): ws3.cell(row=3,column=c,value=h); hdr(ws3,3,len(h3))
for y in range(2019,2027):
    yt=[t for t in closed if t[10]==y]
    if not yt: continue
    ypnl=sum(t[8] for t in yt); yw=len([t for t in yt if t[8]>0]); yl=len([t for t in yt if t[8]<0])
    prev=[t for t in closed if t[10]<y]; ys=prev[-1][9] if prev else INIT; ye=yt[-1][9]; yr_=round((ye-ys)/ys*100,2) if ys>0 else 0
    vals=[y,len(yt),yw,yl,round(ypnl,2),yr_]
    r=4+y-2019
    for c,v in enumerate(vals,1): ws3.cell(row=r,column=c,value=v); ws3.cell(row=r,column=c).font=dfont; ws3.cell(row=r,column=c).border=bdr; ws3.cell(row=r,column=c).alignment=Alignment(horizontal="center")
    if yr_>0: ws3.cell(row=r,column=5).font=gf; ws3.cell(row=r,column=5).fill=gfill; ws3.cell(row=r,column=6).font=gf; ws3.cell(row=r,column=6).fill=gfill
    elif yr_<0: ws3.cell(row=r,column=5).font=rf_; ws3.cell(row=r,column=5).fill=rfill; ws3.cell(row=r,column=6).font=rf_; ws3.cell(row=r,column=6).fill=rfill
tr_=4+len([y for y in range(2019,2027) if any(t for t in closed if t[10]==y)])+1
# Summary row
ws3.cell(row=tr_,column=1,value="Total").font=Font(name="Microsoft YaHei",size=10,bold=True)
ws3.cell(row=tr_,column=2,value=len(closed)).font=dfont
ws3.cell(row=tr_,column=3,value=len(wins)).font=dfont
ws3.cell(row=tr_,column=4,value=len(losses)).font=dfont
ws3.cell(row=tr_,column=5,value=round(final_bal-INIT,2)).font=dfont
ws3.cell(row=tr_,column=6,value=total_ret).font=dfont
for c in range(1,7): ws3.column_dimensions[get_column_letter(c)].width=14; ws3.cell(row=tr_,column=c).border=bdr

# Sheet 4: Stats
ws4=wb.create_sheet("Stats")
ws4.merge_cells("A1:C1"); ws4.cell(row=1,column=1,value="Statistics").font=tfont
stats=[
    ("Strategy","Volty EMA50 + RSI + DynamicSize(1.2x)",""),
    ("Config","{}%x{}x SL:{}% TP:{}%".format(PCT,LEV,SL,TP),""),
    ("Start","${}".format(INIT),""), ("Final","${:,.0f}".format(final_bal),""),
    ("Return","{:+.1f}%".format(total_ret),""), ("","",""),
    ("Total Trades",len(trades),""), ("Closed",len(closed),""),
    ("Wins",len(wins),""), ("Losses",len(losses),""),
    ("Win Rate","{:.1f}%".format(wr),""),
    ("Avg Win","${:+,.0f}".format(np.mean([t[8] for t in wins]) if wins else 0),""),
    ("Avg Loss","${:+,.0f}".format(np.mean([t[8] for t in losses]) if losses else 0),""),
    ("Data","{} -> {}".format(str(df.index[0])[:10],str(df.index[-1])[:10]),""),
    ("Candles","{:,}".format(n),""),
]
for i,item in enumerate(stats):
    k,v,n=item[0],item[1],item[2]
    r=i+3
    ws4.cell(row=r,column=1,value=k).font=Font(name="Microsoft YaHei",size=10,bold=True)
    ws4.cell(row=r,column=2,value=v).font=dfont; ws4.cell(row=r,column=3,value=n).font=sf
    for c in range(1,4): ws4.cell(row=r,column=c).border=bdr
ws4.column_dimensions["A"].width=18; ws4.column_dimensions["B"].width=30

out="/opt/trading/full_trade_log_900.xlsx"
wb.save(out)
print("Saved:",out,"  Trades:",len(trades),"  Final: ${:,.0f}".format(final_bal))
