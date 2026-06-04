"""
优化策略完整回测报表
策略: Volty EMA50 + RSI方向确认 + 动态仓位 + TP7%
数据: 币安 ETH/USDT:USDT 15m 真实K线 (228,514根, 2019-2026)
起始: $700  配置: 60%x3x  SL5% TP7%
"""
import pandas as pd, numpy as np, os, platform
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INIT=700; PCT=60; LEV=3; SL=5; TP=7

DATA="/opt/trading/ETHUSDT_15m_full.csv"
print("Loading {}...".format(DATA))
df=pd.read_csv(DATA,parse_dates=["ts"],index_col="ts")
n=len(df)
print("{} candles, {} -> {}".format(n,str(df.index[0])[:10],str(df.index[-1])[:10]))

# Indicators
H=df["high"]; L=df["low"]; C=df["close"]; O=df["open"]
tr=pd.concat([H-L,(H-C.shift(1)).abs(),(L-C.shift(1)).abs()],axis=1).max(axis=1)
atr5=tr.rolling(5).mean()*0.75
ema50=C.ewm(span=50,adjust=False).mean()
atr14=tr.rolling(14).mean()
mid_atr=atr14.median()
d=C.diff(); g=d.clip(lower=0); l=(-d).clip(lower=0)
rsi=100-(100/(1+g.rolling(14).mean()/l.rolling(14).mean()))

# Backtest
print("Backtesting...")
bal=INIT; pos=None; ep=uv=0.0; trades=[]; eq_daily={}
trades_count=0; tp_count=0; sl_count=0

for i in range(200,n):
    hi,lo,op,cl=H.iloc[i],L.iloc[i],O.iloc[i],C.iloc[i]; tm=df.index[i]
    # SL
    if pos=="LONG":
        sl_p=ep*(1-SL/100)
        if lo<=sl_p:
            pnl=(sl_p-ep)/ep*uv; bal+=pnl; sl_count+=1
            trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"LONG","止损",round(ep,2),round(sl_p,2),round(uv,2),round(pnl,2),round(bal,2),tm.year])
            pos=None; eq_daily[str(tm.date())]=bal; continue
    if pos=="SHORT":
        sl_p=ep*(1+SL/100)
        if hi>=sl_p:
            pnl=(ep-sl_p)/ep*uv; bal+=pnl; sl_count+=1
            trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"SHORT","止损",round(ep,2),round(sl_p,2),round(uv,2),round(pnl,2),round(bal,2),tm.year])
            pos=None; eq_daily[str(tm.date())]=bal; continue
    # TP
    if pos=="LONG" and hi>=ep*(1+TP/100):
        tp_p=ep*(1+TP/100)
        pnl=(tp_p-ep)/ep*uv; bal+=pnl; tp_count+=1
        trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"LONG","止盈",round(ep,2),round(tp_p,2),round(uv,2),round(pnl,2),round(bal,2),tm.year])
        pos=None; eq_daily[str(tm.date())]=bal; continue
    if pos=="SHORT" and lo<=ep*(1-TP/100):
        tp_p=ep*(1-TP/100)
        pnl=(ep-tp_p)/ep*uv; bal+=pnl; tp_count+=1
        trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"SHORT","止盈",round(ep,2),round(tp_p,2),round(uv,2),round(pnl,2),round(bal,2),tm.year])
        pos=None; eq_daily[str(tm.date())]=bal; continue

    # Volty + RSI signals
    pc=C.iloc[i-1]; pa=atr5.iloc[i-1]; pe=ema50.iloc[i-1]
    if pd.isna(pa) or pa<=0: continue
    ll_v=pc+pa; sl_v=pc-pa
    tb=pc>pe; tbe=pc<pe
    rb=rsi.iloc[i-1]>50 if not pd.isna(rsi.iloc[i-1]) else False
    rbe=rsi.iloc[i-1]<50 if not pd.isna(rsi.iloc[i-1]) else False

    lt=hi>=ll_v and tb and rb and pos!="LONG"
    st=lo<=sl_v and tbe and rbe and pos!="SHORT"

    if lt and st:
        if abs(op-ll_v)>abs(op-sl_v): lt=False
        else: st=False

    if lt:
        if pos=="SHORT":
            xp=ll_v; pl=(ep-xp)/ep*uv; bal+=pl
            trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"SHORT->LONG","翻转",round(ep,2),round(xp,2),round(uv,2),round(pl,2),round(bal,2),tm.year])
            pos=None; eq_daily[str(tm.date())]=bal
        if bal>10:
            cap=bal*PCT/100
            av=atr14.iloc[i] if not pd.isna(atr14.iloc[i]) else mid_atr
            if av>0: cap*=min(1.5,max(0.5,mid_atr/av))
            cap=min(cap,bal*0.98); uv=cap*LEV; ep=ll_v; pos="LONG"; trades_count+=1
            eq_daily[str(tm.date())]=bal
    elif st:
        if pos=="LONG":
            xp=sl_v; pl=(xp-ep)/ep*uv; bal+=pl
            trades.append([len(trades)+1,str(tm.date()),str(tm.time())[:5],"LONG->SHORT","翻转",round(ep,2),round(xp,2),round(uv,2),round(pl,2),round(bal,2),tm.year])
            pos=None; eq_daily[str(tm.date())]=bal
        if bal>10:
            cap=bal*PCT/100
            av=atr14.iloc[i] if not pd.isna(atr14.iloc[i]) else mid_atr
            if av>0: cap*=min(1.5,max(0.5,mid_atr/av))
            cap=min(cap,bal*0.98); uv=cap*LEV; ep=sl_v; pos="SHORT"; trades_count+=1
            eq_daily[str(tm.date())]=bal

if pos:
    lp=C.iloc[-1]; pl=(lp-ep)/ep*uv if pos=="LONG" else (ep-lp)/ep*uv; bal+=pl
    trades.append([len(trades)+1,str(df.index[-1].date()),str(df.index[-1].time())[:5],pos,"最终平仓",round(ep,2),round(lp,2),round(uv,2),round(pl,2),round(bal,2),df.index[-1].year])

final_bal=round(bal,2); total_ret=round((final_bal-INIT)/INIT*100,2)
closed=[t for t in trades if t[8]!=0]
wins=[t for t in closed if t[8]>0]; losses=[t for t in closed if t[8]<0]
wr=len(wins)/len(closed)*100 if closed else 0

# Yearly
yearly=[]
for y in range(2019,2027):
    yt=[t for t in closed if t[10]==y]
    if not yt: continue
    ypnl=sum(t[8] for t in yt)
    prev=[t for t in closed if t[10]<y]
    ys=prev[-1][9] if prev else INIT
    ye=yt[-1][9]
    yearly.append([y,len(yt),len([t for t in yt if t[8]>0]),len([t for t in yt if t[8]<0]),round(ypnl,2),round((ye-ys)/ys*100,2) if ys>0 else 0,round(ys,2),round(ye,2)])

# Monthly
months=[]
for y in range(2019,2027):
    for m in range(1,13):
        mt=[t for t in closed if t[10]==y and int(t[1][5:7])==m]
        label="{}/{:02d}".format(y,m)
        if not mt: months.append([label,0,0,0,0,0,0,0]); continue
        mpnl=sum(t[8] for t in mt)
        prev=[t for t in closed if t[10]<y or (t[10]==y and int(t[1][5:7])<m)]
        ms=prev[-1][9] if prev else INIT
        me=mt[-1][9]
        mr=round((me-ms)/ms*100,2) if ms>0 else 0
        months.append([label,len(mt),len([t for t in mt if t[8]>0]),len([t for t in mt if t[8]<0]),round(mpnl,2),mr,round(ms,2),round(me,2)])

# DD
eq_df=pd.DataFrame(list(eq_daily.items()),columns=["date","eq"])
eq_df["date"]=pd.to_datetime(eq_df["date"]); eq_df=eq_df.sort_values("date"); eq_df.set_index("date",inplace=True)
eq_df["peak"]=eq_df["eq"].cummax(); eq_df["dd"]=(eq_df["eq"]-eq_df["peak"])/eq_df["peak"]*100
max_dd=round(eq_df["dd"].min(),2); max_dd_date=str(eq_df["dd"].idxmin())[:10]

# === Excel ===
print("Generating Excel...")
wb=Workbook()
hfont=Font(name="Microsoft YaHei",size=11,bold=True,color="FFFFFF")
hfill=PatternFill(start_color="1a1a2e",end_color="1a1a2e",fill_type="solid")
tfont=Font(name="Microsoft YaHei",size=16,bold=True,color="1a1a2e")
sf=Font(name="Microsoft YaHei",size=10,color="666666")
dfont=Font(name="Microsoft YaHei",size=10)
gf=Font(name="Microsoft YaHei",size=10,color="006100"); rf_=Font(name="Microsoft YaHei",size=10,color="9C0006")
gfill=PatternFill(start_color="C6EFCE",end_color="C6EFCE",fill_type="solid")
rfill=PatternFill(start_color="FFC7CE",end_color="FFC7CE",fill_type="solid")
yfill=PatternFill(start_color="FFEB9C",end_color="FFEB9C",fill_type="solid")
lfill=PatternFill(start_color="F5F5F5",end_color="F5F5F5",fill_type="solid")
bdr=Border(left=Side(style="thin",color="D0D0D0"),right=Side(style="thin",color="D0D0D0"),top=Side(style="thin",color="D0D0D0"),bottom=Side(style="thin",color="D0D0D0"))

def hdr(ws,row,cols):
    for c in range(1,cols+1): cl=ws.cell(row=row,column=c); cl.font=hfont; cl.fill=hfill; cl.alignment=Alignment(horizontal="center",vertical="center"); cl.border=bdr

def sdr(ws,row,cols,even=False):
    for c in range(1,cols+1): cl=ws.cell(row=row,column=c); cl.font=dfont; cl.border=bdr; cl.alignment=Alignment(horizontal="center")
    if even:
        for c in range(1,cols+1): ws.cell(row=row,column=c).fill=lfill

# Sheet 1
ws1=wb.active; ws1.title="每笔交易"
ws1.merge_cells("A1:K1")
ws1.cell(row=1,column=1,value="优化策略回测 Volty+RSI+动态仓位+TP7%").font=tfont
ws1.merge_cells("A2:K2")
ws1.cell(row=2,column=1,value="币安ETHUSDT 15m真实K线 2019-2026 | ${}起 {}%x{}x SL{}% TP{}%".format(INIT,PCT,LEV,SL,TP)).font=sf
h1=["序号","日期","时间","方向","动作","入场价","出场价","持仓价值","盈亏","余额","年份"]
for c,h in enumerate(h1,1): ws1.cell(row=4,column=c,value=h)
hdr(ws1,4,len(h1))
for i in range(len(trades)):
    r=i+5; t=trades[i]
    for c in range(len(h1)): ws1.cell(row=r,column=c+1,value=t[c] if c<10 else t[c])
    sdr(ws1,r,len(h1),i%2==0)
    pc_=ws1.cell(row=r,column=9)
    if t[8]>0: pc_.font=gf; pc_.fill=gfill
    elif t[8]<0: pc_.font=rf_; pc_.fill=rfill
ws1.freeze_panes="A5"
for c in range(1,len(h1)+1): ws1.column_dimensions[get_column_letter(c)].width=14

# Sheet 2: Yearly
ws2=wb.create_sheet("年度汇总")
ws2.merge_cells("A1:H1"); ws2.cell(row=1,column=1,value="年度收益汇总").font=tfont
h2=["年份","交易次数","盈利次数","亏损次数","盈亏","收益率%","年初余额","年末余额"]
for c,h in enumerate(h2,1): ws2.cell(row=3,column=c,value=h)
hdr(ws2,3,len(h2))
for i,yr in enumerate(yearly):
    r=i+4
    for c in range(len(h2)): ws2.cell(row=r,column=c+1,value=yr[c])
    sdr(ws2,r,len(h2),i%2==0)
    rc_=ws2.cell(row=r,column=6)
    if yr[5]>0: rc_.font=gf; rc_.fill=gfill
    elif yr[5]<0: rc_.font=rf_; rc_.fill=rfill
sr=len(yearly)+5
ws2.cell(row=sr,column=1,value="合计").font=Font(name="Microsoft YaHei",size=10,bold=True)
for c in range(1,5): ws2.cell(row=sr,column=c+1,value=sum(yr[c] for yr in yearly))
ws2.cell(row=sr,column=5,value=round(sum(yr[4] for yr in yearly),2))
ws2.cell(row=sr,column=6,value=total_ret); ws2.cell(row=sr,column=7,value=INIT); ws2.cell(row=sr,column=8,value=final_bal)
for c in range(1,9): ws2.cell(row=sr,column=c).font=Font(name="Microsoft YaHei",size=10,bold=True); ws2.cell(row=sr,column=c).border=bdr; ws2.cell(row=sr,column=c).fill=yfill
for c in range(1,9): ws2.column_dimensions[get_column_letter(c)].width=16
ws2.freeze_panes="A4"

# Sheet 3: Monthly
ws3=wb.create_sheet("月度汇总")
ws3.merge_cells("A1:H1"); ws3.cell(row=1,column=1,value="月度收益汇总").font=tfont
h3=["月份","交易次数","盈利次数","亏损次数","盈亏","收益率%","月初余额","月末余额"]
for c,h in enumerate(h3,1): ws3.cell(row=3,column=c,value=h)
hdr(ws3,3,len(h3))
for i,mo in enumerate(months):
    r=i+4
    for c in range(len(h3)): ws3.cell(row=r,column=c+1,value=mo[c])
    sdr(ws3,r,len(h3),i%2==0)
    rc_=ws3.cell(row=r,column=6)
    if mo[5]>0: rc_.font=gf; rc_.fill=gfill
    elif mo[5]<0: rc_.font=rf_; rc_.fill=rfill
for c in range(1,9): ws3.column_dimensions[get_column_letter(c)].width=16
ws3.freeze_panes="A4"

# Sheet 4: Statistics
ws4=wb.create_sheet("统计")
ws4.merge_cells("A1:C1"); ws4.cell(row=1,column=1,value="策略统计").font=tfont
stats=[
    ("策略","Volty EMA50 + RSI方向 + 动态仓位",""),
    ("标的","ETH/USDT:USDT 永续合约",""),
    ("K线","15分钟",""),
    ("数据范围","{} -> {}".format(str(df.index[0])[:10],str(df.index[-1])[:10]),""),
    ("总K线数","{:,}".format(n),""),
    ("","",""),
    ("配置","{}% x {}x | SL:{}% | TP:{}%".format(PCT,LEV,SL,TP),""),
    ("起始资金","${:,.2f}".format(INIT),""),
    ("最终资金","${:,.2f}".format(final_bal),""),
    ("总收益","${:+,.2f}".format(final_bal-INIT),""),
    ("总收益率","{:+.2f}%".format(total_ret),""),
    ("","",""),
    ("总成交笔数",len(trades),""),
    ("平仓次数",len(closed),""),
    ("止盈次数",tp_count,""),
    ("止损次数",sl_count,""),
    ("盈利次数",len(wins),""),
    ("亏损次数",len(losses),""),
    ("胜率","{:.1f}%".format(wr),""),
    ("平均盈利","${:+,.0f}".format(np.mean([t[8] for t in wins]) if wins else 0),""),
    ("平均亏损","${:+,.0f}".format(np.mean([t[8] for t in losses]) if losses else 0),""),
    ("盈亏比","{:.2f}".format(abs(np.mean([t[8] for t in wins])/np.mean([t[8] for t in losses])) if wins and losses else 0),""),
    ("最大回撤","{:+.2f}%".format(max_dd),max_dd_date),
    ("","",""),
    ("盈利年份","{}/{}".format(len([y for y in yearly if y[5]>0]),len(yearly)),""),
    ("最佳年","{}年 ({:+.1f}%)".format(max(yearly,key=lambda x:x[5])[0],max(yearly,key=lambda x:x[5])[5]),""),
    ("最差年","{}年 ({:+.1f}%)".format(min(yearly,key=lambda x:x[5])[0],min(yearly,key=lambda x:x[5])[5]),""),
]
for i,(k,v,note) in enumerate(stats):
    r=i+3
    ws4.cell(row=r,column=1,value=k).font=Font(name="Microsoft YaHei",size=10,bold=True)
    ws4.cell(row=r,column=2,value=v).font=dfont; ws4.cell(row=r,column=3,value=note).font=sf
    for c in range(1,4): ws4.cell(row=r,column=c).border=bdr
    if i%2:
        for c in range(1,4): ws4.cell(row=r,column=c).fill=lfill
ws4.column_dimensions["A"].width=18; ws4.column_dimensions["B"].width=30; ws4.column_dimensions["C"].width=15

# Save
out="/opt/trading/optimized_report.xlsx"
wb.save(out)
print("\nSaved: {}".format(out))
print("Final: ${:,.0f} ({:+.1f}%) | {} trades | WR: {:.1f}% | MaxDD: {:+.1f}%".format(final_bal,total_ret,len(closed),wr,max_dd))
for y in yearly: print("  {}: ${:,.0f} -> ${:,.0f} ({:+.1f}%)".format(y[0],y[6],y[7],y[5]))
