"""
2025年 Volty 策略完整回测 → Excel 报表
仓位: 50% × 3x | 起始: $700
"""
import ccxt, pandas as pd, numpy as np, time as _time
from datetime import datetime, timezone, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
import os

VLEN, VMULT, PCT, LEV, SL_PCT = 5, 0.75, 50, 3, 5
INIT = 700.0
SYMBOL = "ETH/USDT:USDT"
TIMEFRAME = "15m"

print("=" * 60)
print("  2025年 Volty 策略完整回测 → Excel 报表")
print(f"  仓位: {PCT}% × {LEV}x | 起始: ${INIT}")
print("=" * 60)

# ═══════════════ 下载数据 ═══════════════
print("\n>>> 下载 2025年 ETHUSDT 15m 数据...")
ex = ccxt.binance({"enableRateLimit": True, "timeout": 30000})
since = int(datetime(2025, 1, 1, tzinfo=timezone.utc).timestamp() * 1000)
end_ts = int(datetime(2026, 1, 1, tzinfo=timezone.utc).timestamp() * 1000)

all_klines = []
while since < end_ts:
    try:
        klines = ex.fetch_ohlcv(SYMBOL, TIMEFRAME, since=since, limit=1000)
        if not klines: break
        all_klines.extend(klines)
        since = klines[-1][0] + 1
        if len(all_klines) % 5000 == 0:
            print(f"  {len(all_klines)} 根...")
    except Exception as e:
        _time.sleep(1)

df = pd.DataFrame(all_klines, columns=["ts", "open", "high", "low", "close", "volume"])
df["ts"] = pd.to_datetime(df["ts"], unit="ms", utc=True)
df.set_index("ts", inplace=True)
df = df[(df.index >= "2025-01-01") & (df.index < "2026-01-01")]
print(f"  总计: {len(df)} 根K线")

# ═══════════════ 预计算 ═══════════════
tr = pd.concat([df["high"] - df["low"], (df["high"] - df["close"].shift(1)).abs(), (df["low"] - df["close"].shift(1)).abs()], axis=1).max(axis=1)
atr_sig = tr.rolling(VLEN).mean() * VMULT
ema50 = df["close"].ewm(span=50, adjust=False).mean()

# ═══════════════ 回测 ═══════════════
print(">>> 回测中...")
bal = INIT
pos = None; ep = cts = uv = 0.0
trades = []
equity_daily = {}  # date -> equity
START = max(VLEN + 3, 53)

for i in range(START, len(df)):
    hi, lo, op = df["high"].iloc[i], df["low"].iloc[i], df["open"].iloc[i]
    tm = df.index[i]
    date_key = str(tm.date())

    # Stop loss
    if pos == "LONG":
        sp = ep * (1 - SL_PCT / 100)
        if lo <= sp:
            pnl = (sp - ep) / ep * uv
            bal += pnl
            trades.append({
                "序号": len(trades) + 1, "日期": str(tm.date()), "时间": str(tm.time())[:5],
                "方向": "做多", "动作": "止损平仓",
                "入场价": round(ep, 2), "出场价": round(sp, 2),
                "持仓价值U": round(uv, 2), "盈亏U": round(pnl, 2), "余额U": round(bal, 2),
                "周": tm.isocalendar()[1], "月": tm.month
            })
            pos = None; equity_daily[date_key] = bal; continue

    if pos == "SHORT":
        sp = ep * (1 + SL_PCT / 100)
        if hi >= sp:
            pnl = (ep - sp) / ep * uv
            bal += pnl
            trades.append({
                "序号": len(trades) + 1, "日期": str(tm.date()), "时间": str(tm.time())[:5],
                "方向": "做空", "动作": "止损平仓",
                "入场价": round(ep, 2), "出场价": round(sp, 2),
                "持仓价值U": round(uv, 2), "盈亏U": round(pnl, 2), "余额U": round(bal, 2),
                "周": tm.isocalendar()[1], "月": tm.month
            })
            pos = None; equity_daily[date_key] = bal; continue

    pa = atr_sig.iloc[i - 1]
    if pd.isna(pa) or pa <= 0: continue
    pc = df["close"].iloc[i - 1]
    ll, sl = pc + pa, pc - pa

    # Trend
    pe = ema50.iloc[i - 1]
    tb = df["close"].iloc[i - 1] > pe
    tbe = df["close"].iloc[i - 1] < pe

    lt = hi >= ll and tb and pos != "LONG"
    st = lo <= sl and tbe and pos != "SHORT"
    if lt and st:
        if abs(op - ll) > abs(op - sl): lt = False
        else: st = False

    if lt:
        if pos == "SHORT":
            xp = ll; pl = (ep - xp) / ep * uv; bal += pl
            trades.append({
                "序号": len(trades) + 1, "日期": str(tm.date()), "时间": str(tm.time())[:5],
                "方向": "做空→做多", "动作": "翻转平仓",
                "入场价": round(ep, 2), "出场价": round(xp, 2),
                "持仓价值U": round(uv, 2), "盈亏U": round(pl, 2), "余额U": round(bal, 2),
                "周": tm.isocalendar()[1], "月": tm.month
            })
            pos = None; equity_daily[date_key] = bal
        if bal > 10:
            cap = bal * PCT / 100.0; uv = cap * LEV; ep = ll; cts = uv / ep; pos = "LONG"
            trades.append({
                "序号": len(trades) + 1, "日期": str(tm.date()), "时间": str(tm.time())[:5],
                "方向": "做多", "动作": "开仓",
                "入场价": round(ep, 2), "出场价": 0,
                "持仓价值U": round(uv, 2), "盈亏U": 0, "余额U": round(bal, 2),
                "周": tm.isocalendar()[1], "月": tm.month
            })
            equity_daily[date_key] = bal

    elif st:
        if pos == "LONG":
            xp = sl; pl = (xp - ep) / ep * uv; bal += pl
            trades.append({
                "序号": len(trades) + 1, "日期": str(tm.date()), "时间": str(tm.time())[:5],
                "方向": "做多→做空", "动作": "翻转平仓",
                "入场价": round(ep, 2), "出场价": round(xp, 2),
                "持仓价值U": round(uv, 2), "盈亏U": round(pl, 2), "余额U": round(bal, 2),
                "周": tm.isocalendar()[1], "月": tm.month
            })
            pos = None; equity_daily[date_key] = bal
        if bal > 10:
            cap = bal * PCT / 100.0; uv = cap * LEV; ep = sl; cts = uv / ep; pos = "SHORT"
            trades.append({
                "序号": len(trades) + 1, "日期": str(tm.date()), "时间": str(tm.time())[:5],
                "方向": "做空", "动作": "开仓",
                "入场价": round(ep, 2), "出场价": 0,
                "持仓价值U": round(uv, 2), "盈亏U": 0, "余额U": round(bal, 2),
                "周": tm.isocalendar()[1], "月": tm.month
            })
            equity_daily[date_key] = bal

# Final close
if pos:
    lp = df["close"].iloc[-1]
    pl = (lp - ep) / ep * uv if pos == "LONG" else (ep - lp) / ep * uv
    bal += pl
    trades.append({
        "序号": len(trades) + 1, "日期": str(df.index[-1].date()), "时间": str(df.index[-1].time())[:5],
        "方向": "做多" if pos == "LONG" else "做空", "动作": "最终平仓",
        "入场价": round(ep, 2), "出场价": round(lp, 2),
        "持仓价值U": round(uv, 2), "盈亏U": round(pl, 2), "余额U": round(bal, 2),
        "周": df.index[-1].isocalendar()[1], "月": df.index[-1].month
    })

trades_df = pd.DataFrame(trades)
final_bal = round(bal, 2)
total_ret = round((final_bal - INIT) / INIT * 100, 2)

# ═══════════════ 汇总统计 ═══════════════
print(">>> 生成统计...")

# 日权益曲线
eq_df = pd.DataFrame(list(equity_daily.items()), columns=["日期", "权益"])
eq_df["日期"] = pd.to_datetime(eq_df["日期"])
eq_df = eq_df.sort_values("日期")
eq_df.set_index("日期", inplace=True)
eq_df["峰值"] = eq_df["权益"].cummax()
eq_df["回撤"] = (eq_df["权益"] - eq_df["峰值"]) / eq_df["峰值"] * 100

# 收盘交易
closed = trades_df[trades_df.iloc[:, 8] != 0].copy()
wins = closed[closed.iloc[:, 8] > 0]
losses = closed[closed.iloc[:, 8] < 0]

# ── 周度汇总 ──
weekly_data = []
for w in sorted(trades_df.iloc[:, 10].unique()):
    wt = closed[closed.iloc[:, 10] == w]
    if wt.empty: continue
    weekly_pnl = round(wt.iloc[:, 8].sum(), 2)
    w_trades = len(wt)
    w_wins = len(wt[wt.iloc[:, 8] > 0])
    w_last = wt.iloc[-1, 9]
    prev_weeks = closed[closed.iloc[:, 10] < w]
    w_start = prev_weeks.iloc[-1, 9] if not prev_weeks.empty else INIT
    w_ret = round((w_last - w_start) / w_start * 100, 2) if w_start > 0 else 0
    weekly_data.append({
        "周": f"第{w}周", "交易次数": w_trades, "盈利次数": w_wins,
        "亏损次数": w_trades - w_wins, "盈亏U": weekly_pnl, "收益率%": w_ret,
        "周初余额U": round(w_start, 2), "周末余额U": round(w_last, 2)
    })

weekly_df = pd.DataFrame(weekly_data)

# ── 月度汇总 ──
monthly_data = []
for m in range(1, 13):
    mt = closed[closed.iloc[:, 11] == m]
    if mt.empty:
        monthly_data.append({"月份": f"{m}月", "交易次数": 0, "盈利次数": 0, "亏损次数": 0, "盈亏U": 0, "收益率%": 0, "月初余额U": 0, "月末余额U": 0})
        continue
    m_pnl = round(mt.iloc[:, 8].sum(), 2)
    m_trades = len(mt)
    m_wins = len(mt[mt.iloc[:, 8] > 0])
    prev_months = closed[closed.iloc[:, 11] < m]
    m_start = prev_months.iloc[-1, 9] if not prev_months.empty else INIT
    m_end = mt.iloc[-1, 9]
    m_ret = round((m_end - m_start) / m_start * 100, 2) if m_start > 0 else 0
    monthly_data.append({
        "月份": f"{m}月", "交易次数": m_trades, "盈利次数": m_wins,
        "亏损次数": m_trades - m_wins, "盈亏U": m_pnl, "收益率%": m_ret,
        "月初余额U": round(m_start, 2), "月末余额U": round(m_end, 2)
    })

monthly_df = pd.DataFrame(monthly_data)

# ── 回撤分析 ──
max_dd = round(eq_df["回撤"].min(), 2)
max_dd_date = str(eq_df["回撤"].idxmin())[:10]
avg_dd = round(eq_df["回撤"].mean(), 2)

# ═══════════════ 生成 Excel ═══════════════
print(">>> 生成 Excel...")

wb = Workbook()

# ── 样式定义 ──
header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
title_font = Font(name="微软雅黑", size=16, bold=True, color="1a1a2e")
subtitle_font = Font(name="微软雅黑", size=10, color="666666")
data_font = Font(name="微软雅黑", size=10)
green_font = Font(name="微软雅黑", size=10, color="006100")
red_font = Font(name="微软雅黑", size=10, color="9C0006")
green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
light_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
thin_border = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)

def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

def style_data_row(ws, row, cols, is_even=False):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c > 2 else "left")
        if is_even:
            cell.fill = light_fill

def auto_width(ws, cols):
    for c in range(1, cols + 1):
        max_len = 0
        for row in ws.iter_rows(min_col=c, max_col=c, values_only=True):
            for val in row:
                if val:
                    max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(c)].width = min(max_len + 4, 25)

# ═══ Sheet 1: 交易明细 ═══
ws1 = wb.active
ws1.title = "交易明细"

# Title
ws1.merge_cells("A1:K1")
ws1.cell(row=1, column=1, value="2025年 Volty Expan Close 策略 — 每笔交易明细").font = title_font
ws1.merge_cells("A2:K2")
ws1.cell(row=2, column=1, value=f"仓位: {PCT}% × {LEV}x杠杆 | 起始资金: ${INIT} | 最终资金: ${final_bal:,.0f} | 总收益率: {total_ret:+.1f}%").font = subtitle_font

headers1 = ["序号", "日期", "时间", "方向", "动作", "入场价(USDT)", "出场价(USDT)", "持仓价值(U)", "盈亏(U)", "余额(U)", "周"]
for c, h in enumerate(headers1, 1):
    ws1.cell(row=4, column=c, value=h)
style_header(ws1, 4, len(headers1))

for i in range(len(trades_df)):
    r = i + 5
    row = trades_df.iloc[i]
    vals = [int(row.iloc[0]), str(row.iloc[1]), str(row.iloc[2]), str(row.iloc[3]), str(row.iloc[4]),
            float(row.iloc[5]), float(row.iloc[6]), float(row.iloc[7]), float(row.iloc[8]), float(row.iloc[9]),
            f"第{int(row.iloc[10])}周"]
    for c, v in enumerate(vals, 1):
        ws1.cell(row=r, column=c, value=v)
    style_data_row(ws1, r, len(headers1), i % 2 == 0)
    # Color PnL
    pnl_cell = ws1.cell(row=r, column=9)
    if float(row.iloc[8]) > 0:
        pnl_cell.font = green_font; pnl_cell.fill = green_fill
    elif float(row.iloc[8]) < 0:
        pnl_cell.font = red_font; pnl_cell.fill = red_fill

auto_width(ws1, len(headers1))
ws1.auto_filter.ref = f"A4:K{len(trades_df)+4}"
ws1.freeze_panes = "A5"

# ═══ Sheet 2: 月度汇总 ═══
ws2 = wb.create_sheet("月度汇总")
ws2.merge_cells("A1:H1")
ws2.cell(row=1, column=1, value="2025年 月度收益汇总").font = title_font

headers2 = ["月份", "交易次数", "盈利次数", "亏损次数", "盈亏(U)", "收益率%", "月初余额(U)", "月末余额(U)"]
for c, h in enumerate(headers2, 1):
    ws2.cell(row=3, column=c, value=h)
style_header(ws2, 3, len(headers2))

for i, (_, row) in enumerate(monthly_df.iterrows()):
    r = i + 4
    ws2.cell(row=r, column=1, value=row.iloc[0])
    ws2.cell(row=r, column=2, value=row.iloc[1])
    ws2.cell(row=r, column=3, value=row.iloc[2])
    ws2.cell(row=r, column=4, value=row.iloc[3])
    ws2.cell(row=r, column=5, value=row.iloc[4])
    ws2.cell(row=r, column=6, value=row.iloc[5])
    ws2.cell(row=r, column=7, value=row.iloc[6])
    ws2.cell(row=r, column=8, value=row.iloc[7])
    style_data_row(ws2, r, len(headers2), i % 2 == 0)
    # Color return
    ret_cell = ws2.cell(row=r, column=6)
    if row.iloc[5] > 0:
        ret_cell.font = green_font; ret_cell.fill = green_fill
    elif row.iloc[5] < 0:
        ret_cell.font = red_font; ret_cell.fill = red_fill

# Summary row
sr = len(monthly_df) + 5
ws2.cell(row=sr, column=1, value="全年合计").font = Font(name="微软雅黑", size=10, bold=True)
ws2.cell(row=sr, column=2, value=monthly_df.iloc[:, 1].sum())
ws2.cell(row=sr, column=3, value=monthly_df.iloc[:, 2].sum())
ws2.cell(row=sr, column=4, value=monthly_df.iloc[:, 3].sum())
ws2.cell(row=sr, column=5, value=round(monthly_df.iloc[:, 4].sum(), 2))
ws2.cell(row=sr, column=6, value=total_ret)
ws2.cell(row=sr, column=7, value=INIT)
ws2.cell(row=sr, column=8, value=final_bal)
for c in range(1, len(headers2) + 1):
    ws2.cell(row=sr, column=c).font = Font(name="微软雅黑", size=10, bold=True)
    ws2.cell(row=sr, column=c).border = thin_border
    ws2.cell(row=sr, column=c).fill = yellow_fill

auto_width(ws2, len(headers2))
ws2.freeze_panes = "A4"

# ═══ Sheet 3: 周度汇总 ═══
ws3 = wb.create_sheet("周度汇总")
ws3.merge_cells("A1:H1")
ws3.cell(row=1, column=1, value="2025年 周度收益汇总").font = title_font

headers3 = ["周", "交易次数", "盈利次数", "亏损次数", "盈亏(U)", "收益率%", "周初余额(U)", "周末余额(U)"]
for c, h in enumerate(headers3, 1):
    ws3.cell(row=3, column=c, value=h)
style_header(ws3, 3, len(headers3))

for i in range(len(weekly_df)):
    r = i + 4
    row = weekly_df.iloc[i]
    ws3.cell(row=r, column=1, value=str(row.iloc[0]))
    ws3.cell(row=r, column=2, value=int(row.iloc[1]))
    ws3.cell(row=r, column=3, value=int(row.iloc[2]))
    ws3.cell(row=r, column=4, value=int(row.iloc[3]))
    ws3.cell(row=r, column=5, value=float(row.iloc[4]))
    ws3.cell(row=r, column=6, value=float(row.iloc[5]))
    ws3.cell(row=r, column=7, value=float(row.iloc[6]))
    ws3.cell(row=r, column=8, value=float(row.iloc[7]))
    style_data_row(ws3, r, len(headers3), i % 2 == 0)
    ret_cell = ws3.cell(row=r, column=6)
    if float(row.iloc[5]) > 0:
        ret_cell.font = green_font; ret_cell.fill = green_fill
    elif float(row.iloc[5]) < 0:
        ret_cell.font = red_font; ret_cell.fill = red_fill

auto_width(ws3, len(headers3))
ws3.freeze_panes = "A4"

# ═══ Sheet 4: 总体统计 ═══
ws4 = wb.create_sheet("总体统计")
ws4.merge_cells("A1:B1")
ws4.cell(row=1, column=1, value="2025年 策略总体统计").font = title_font

stats = [
    ("策略名称", "Volty Expan Close"),
    ("交易标的", "ETHUSDT 永续合约"),
    ("K线周期", "15分钟"),
    ("策略参数", f"VLEN={VLEN}, MULT={VMULT}"),
    ("趋势过滤", "EMA50"),
    ("仓位配置", f"{PCT}% × {LEV}x杠杆"),
    ("止损比例", f"{SL_PCT}%"),
    ("", ""),
    ("起始资金", f"${INIT:,.2f}"),
    ("最终资金", f"${final_bal:,.2f}"),
    ("总收益", f"${final_bal-INIT:+,.2f}"),
    ("总收益率", f"{total_ret:+.2f}%"),
    ("", ""),
    ("总成交笔数", len(trades_df)),
    ("平仓次数", len(closed)),
    ("盈利次数", len(wins)),
    ("亏损次数", len(losses)),
    ("胜率", f"{len(wins)/len(closed)*100:.1f}%" if len(closed) > 0 else "0%"),
    ("平均盈利", f"${wins.iloc[:, 8].mean():+,.0f}" if len(wins) > 0 else "N/A"),
    ("平均亏损", f"${losses.iloc[:, 8].mean():+,.0f}" if len(losses) > 0 else "N/A"),
    ("盈亏比", f"{abs(wins.iloc[:, 8].mean()/losses.iloc[:, 8].mean()):.2f}" if len(wins) > 0 and len(losses) > 0 else "N/A"),
    ("最大单笔盈利", f"${wins.iloc[:, 8].max():+,.0f}" if len(wins) > 0 else "N/A"),
    ("最大单笔亏损", f"${losses.iloc[:, 8].min():+,.0f}" if len(losses) > 0 else "N/A"),
    ("", ""),
    ("最大回撤", f"{max_dd:+.2f}%"),
    ("最大回撤日期", max_dd_date),
    ("平均回撤", f"{avg_dd:+.2f}%"),
    ("", ""),
    ("盈利月份", f"{len([m for m in monthly_data if m['收益率%'] > 0])}/12"),
    ("盈利周数", f"{len([w for w in weekly_data if w['收益率%'] > 0])}/{len(weekly_data)}"),
    ("最佳月份", f"{max(monthly_data, key=lambda x: x['收益率%'])['月份']} ({max(monthly_data, key=lambda x: x['收益率%'])['收益率%']:+.1f}%)"),
    ("最差月份", f"{min(monthly_data, key=lambda x: x['收益率%'])['月份']} ({min(monthly_data, key=lambda x: x['收益率%'])['收益率%']:+.1f}%)"),
]

for i, (label, value) in enumerate(stats):
    r = i + 3
    ws4.cell(row=r, column=1, value=label).font = Font(name="微软雅黑", size=10, bold=True)
    ws4.cell(row=r, column=2, value=value).font = data_font
    ws4.cell(row=r, column=1).border = thin_border
    ws4.cell(row=r, column=2).border = thin_border
    if i % 2 == 1:
        ws4.cell(row=r, column=1).fill = light_fill
        ws4.cell(row=r, column=2).fill = light_fill

ws4.column_dimensions["A"].width = 22
ws4.column_dimensions["B"].width = 30

# ═══════════════ 保存 ═══════════════
# Auto-detect platform for output path
import platform
if platform.system() == "Windows":
    output_path = r"C:\Users\周勇\Desktop\量化交易\回测结果\2025年Volty策略回测报表_50pct_3x.xlsx"
else:
    output_path = "/opt/trading/2025年Volty策略回测报表_50pct_3x.xlsx"
out_dir = os.path.dirname(output_path)
if out_dir:
    os.makedirs(out_dir, exist_ok=True)
wb.save(output_path)

print(f"\n{'='*60}")
print(f"  ✅ Excel 报表已保存:")
print(f"  {output_path}")
print(f"{'='*60}")
print(f"")
print(f"  📊 最终结果: ${INIT:,.0f} → ${final_bal:,.0f} ({total_ret:+.1f}%)")
print(f"  📊 共 {len(trades_df)} 笔交易，{len(closed)} 次平仓")
print(f"  📊 {len(wins)}赢/{len(losses)}输，胜率 {len(wins)/len(closed)*100:.1f}%")
print(f"  📊 最大回撤: {max_dd:+.1f}%")
print(f"  📊 最佳月: {max(monthly_data, key=lambda x: x['收益率%'])['月份']}，最差月: {min(monthly_data, key=lambda x: x['收益率%'])['月份']}")
print(f"")
print(f"  报表包含 4 个工作表:")
print(f"    1. 交易明细 — 全部 {len(trades_df)} 笔交易")
print(f"    2. 月度汇总 — 12个月统计")
print(f"    3. 周度汇总 — {len(weekly_data)} 周统计")
print(f"    4. 总体统计 — 策略指标")
