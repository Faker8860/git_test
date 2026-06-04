import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INIT=900; PCT=60; LEV=3; SL=5; TP=7

df=pd.read_csv("/opt/trading/ETHUSDT_15m_full.csv",parse_dates=["ts"],index_col="ts")
n=len(df)
H=df["high"].values; L=df["low"].values; C=df["close"].values; O=df["open"].values
tr=pd.concat([df["high"]-df["low"],(df["high"]-df["close"].shift(1)).abs(),(df["low"]-df["close"].shift(1)).abs()],axis=1).max(axis=1)
atr5_v=tr.rolling(5).mean()*0.75; ema50_v=df["close"].ewm(span=50,adjust=False).mean()
atr14_v=tr.rolling(14).mean(); mid_atr=atr14_v.median()
d=df["close"].diff(); g=d.clip(lower=0); l=(-d).clip(lower=0)
rsi_v=(100-(100/(1+g.rolling(14).mean()/l.rolling(14).mean()))).values

bal=INIT; pos=None; ep=uv=0.0; eq={}
for i in range(200,n):
    hi,lo,cl=H[i],L[i],C[i]; tm=df.index[i]
    if pos==1 and lo<=ep*0.95: bal+=(ep*0.95-ep)/ep*uv; pos=None; continue
    if pos==-1 and hi>=ep*1.05: bal+=(ep-ep*1.05)/ep*uv; pos=None; continue
    if pos==1 and hi>=ep*1.07: bal+=(ep*1.07-ep)/ep*uv; pos=None; continue
    if pos==-1 and lo<=ep*0.93: bal+=(ep-ep*0.93)/ep*uv; pos=None; continue
    pc=C[i-1]; pa=atr5_v.iloc[i-1]; pe=ema50_v.iloc[i-1]
    if pd.isna(pa) or pa<=0: continue
    rb=rsi_v[i-1]>50 if not pd.isna(rsi_v[i-1]) else True
    rbe=rsi_v[i-1]<50 if not pd.isna(rsi_v[i-1]) else True
    lt=hi>=pc+pa and pc>pe and rb and pos!=1
    st=lo<=pc-pa and pc<pe and rbe and pos!=-1
    if lt and st:
        if abs(O[i]-(pc+pa))>abs(O[i]-(pc-pa)): lt=False
        else: st=False
    if lt:
        if pos==-1: bal+=(ep-(pc+pa))/ep*uv; pos=None
        if bal>10:
            cap=bal*PCT/100; av=atr14_v.iloc[i] if not pd.isna(atr14_v.iloc[i]) else mid_atr
            if av>0: cap*=min(1.2,max(0.5,mid_atr/av))
            uv=min(cap,bal*0.98)*LEV; ep=pc+pa; pos=1
    elif st:
        if pos==1: bal+=((pc-pa)-ep)/ep*uv; pos=None
        if bal>10:
            cap=bal*PCT/100; av=atr14_v.iloc[i] if not pd.isna(atr14_v.iloc[i]) else mid_atr
            if av>0: cap*=min(1.2,max(0.5,mid_atr/av))
            uv=min(cap,bal*0.98)*LEV; ep=pc-pa; pos=-1
    eq[str(tm.date())]=bal
if pos: bal+=((C[-1]-ep)/ep*uv if pos==1 else (ep-C[-1])/ep*uv)
eq[str(df.index[-1])[:10]]=bal

final_bal=round(bal,2); total_ret=round((final_bal-INIT)/INIT*100,2)

months=[]
prev=INIT
for y in range(2019,2027):
    for m in range(1,13):
        label="{}-{:02d}".format(y,m)
        md=[v for k,v in sorted(eq.items()) if k[:7]==label]
        if not md: months.append([label.replace("-","/"),round(prev,0),round(prev,0),0,0]); continue
        ms=md[0]; me=md[-1]; pnl=round(me-ms,2); ret_m=round((me-ms)/ms*100,2) if ms>0 else 0
        months.append([label.replace("-","/"),round(ms,0),round(me,0),pnl,ret_m]); prev=me

# Excel
wb=Workbook()
hfont=Font(name="Microsoft YaHei",size=11,bold=True,color="FFFFFF")
hfill=PatternFill(start_color="1a1a2e",end_color="1a1a2e",fill_type="solid")
tfont=Font(name="Microsoft YaHei",size=16,bold=True,color="1a1a2e")
dfont=Font(name="Microsoft YaHei",size=10)
gf=Font(name="Microsoft YaHei",size=10,color="006100"); rf_=Font(name="Microsoft YaHei",size=10,color="9C0006")
gfill=PatternFill(start_color="C6EFCE",end_color="C6EFCE",fill_type="solid")
rfill=PatternFill(start_color="FFC7CE",end_color="FFC7CE",fill_type="solid")
bdr=Border(left=Side(style="thin"),right=Side(style="thin"),top=Side(style="thin"),bottom=Side(style="thin"))

ws=wb.active; ws.title="Monthly"
ws.merge_cells("A1:E1"); ws.cell(row=1,column=1,value="Current Strategy {}yr Backtest from ${}".format(round((df.index[-1]-df.index[0]).days/365,1),INIT)).font=tfont
ws.merge_cells("A2:E2"); ws.cell(row=2,column=1,value="Volty+RSI+DynSize(1.2x) {}%x{}x SL{}% TP{}%".format(PCT,LEV,SL,TP)).font=Font(name="Microsoft YaHei",size=10,color="666666")
hds=["Month","Start","End","PnL","Return%"]
for c,h in enumerate(hds,1): ws.cell(row=4,column=c,value=h).font=hfont; ws.cell(row=4,column=c).fill=hfill; ws.cell(row=4,column=c).border=bdr; ws.cell(row=4,column=c).alignment=Alignment(horizontal="center")

for i,m in enumerate(months):
    r=i+5
    for c in range(5): ws.cell(row=r,column=c+1,value=m[c]).font=dfont; ws.cell(row=r,column=c+1).border=bdr; ws.cell(row=r,column=c+1).alignment=Alignment(horizontal="center")
    if m[4]>0: ws.cell(row=r,column=4).font=gf; ws.cell(row=r,column=4).fill=gfill; ws.cell(row=r,column=5).font=gf; ws.cell(row=r,column=5).fill=gfill
    elif m[4]<0: ws.cell(row=r,column=4).font=rf_; ws.cell(row=r,column=4).fill=rfill; ws.cell(row=r,column=5).font=rf_; ws.cell(row=r,column=5).fill=rfill

ws.column_dimensions["A"].width=12
for c in ["B","C","D","E"]: ws.column_dimensions[c].width=14
ws.freeze_panes="A5"

out="/opt/trading/current_strategy_900_monthly.xlsx"
wb.save(out)
print("Start: ${} -> Final: ${:,.0f} ({:+.1f}%)".format(INIT,final_bal,total_ret))
for y in range(2019,2027):
    ym=[m for m in months if m[0][:4]==str(y)]
    if not ym: continue
    ys,ye=ym[0][1],ym[-1][2]; ypnl=sum(m[3] for m in ym)
    print("{}: ${:,.0f} -> ${:,.0f} ({:+.1f}%)".format(y,int(ys),int(ye),round((ye-ys)/ys*100,2) if ys>0 else 0))
print("Saved:",out)
